from io import BytesIO

from django.http import HttpResponse
from docx import Document

from .models import Entry


def build_cv_document(entries, selected_achievement_ids, layout,
                      full_name="Your Name", contact=""):
    doc = Document()
    doc.add_heading(full_name, level=0)
    if contact:
        doc.add_paragraph(contact)

    grouped = {key: [] for key, _ in Entry.CATEGORY_CHOICES}
    for entry in entries:
        grouped[entry.category].append(entry)

    all_selected_achievements = []

    for key, label in Entry.CATEGORY_CHOICES:
        cat_entries = grouped.get(key) or []
        if not cat_entries:
            continue
        doc.add_heading(label, level=1)
        for entry in cat_entries:
            para = doc.add_paragraph()
            heading = entry.name
            if entry.organization:
                heading += f" \u2014 {entry.organization}"
            run = para.add_run(heading)
            run.bold = True
            date_range = entry.date_range()
            if date_range:
                date_run = para.add_run(f"   {date_range}")
                date_run.italic = True
            if entry.description:
                doc.add_paragraph(entry.description)
            achievements = [
                a for a in entry.achievements.all()
                if a.id in selected_achievement_ids
            ]
            all_selected_achievements.extend(achievements)
            if layout == "nested":
                for ach in achievements:
                    doc.add_paragraph(ach.display(), style="List Bullet")

    if layout == "section" and all_selected_achievements:
        doc.add_heading("Achievements", level=1)
        for ach in all_selected_achievements:
            doc.add_paragraph(ach.display(), style="List Bullet")

    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer


def docx_response(buffer, filename="cv.docx"):
    response = HttpResponse(
        buffer.read(),
        content_type=(
            "application/vnd.openxmlformats-officedocument"
            ".wordprocessingml.document"
        ),
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response
